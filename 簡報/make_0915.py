"""產生 撿球機9_15.pptx (白底黑字基本版)。用法: python3 簡報/make_0915.py
實驗數據頁讀同資料夾的 data_0915.json (由 experiments/summarize_runs.py 的結果整理)，沒有就放待補字樣。"""
import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
IMG = HERE / 'img'
FONT = 'Arial'
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x55, 0x55, 0x55)

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(5.625)
BLANK = prs.slide_layouts[6]


def new_slide():
    return prs.slides.add_slide(BLANK)


def textbox(slide, x, y, w, h, lines, size=15, bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6):
    """lines: list of str 或 (str, dict) ；dict 可含 indent, bold, size"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for item in lines:
        text, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        indent = opt.get('indent', 0)
        prefix = ('    – ' if indent else '• ') if opt.get('bullet', True) else ''
        run = p.add_run()
        run.text = prefix + text
        run.font.name = FONT
        run.font.size = Pt(opt.get('size', size - (1 if indent else 0)))
        run.font.bold = opt.get('bold', bold)
        run.font.color.rgb = color
    return tb


def title(slide, text):
    textbox(slide, 0.5, 0.3, 9, 0.7, [(text, {'bullet': False})], size=28, bold=True)


def note(slide, text, x=0.5, y=5.0, w=9, h=0.4):
    textbox(slide, x, y, w, h, [(text, {'bullet': False})], size=11, color=GRAY)


def table(slide, rows, x=0.5, y=1.1, w=9, col_w=None, row_h=0.32, size=12):
    n_rows, n_cols = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(row_h * n_rows))
    tbl = shp.table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for r in range(n_rows):
        tbl.rows[r].height = Inches(row_h)
        for c in range(n_cols):
            cell = tbl.cell(r, c)
            cell.text = ''
            cell.margin_left = cell.margin_right = Inches(0.04)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = str(rows[r][c])
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = (r == 0)
            run.font.color.rgb = BLACK
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE) if r == 0 else RGBColor(0xFF, 0xFF, 0xFF)
    return tbl


def picture(slide, name, x, y, w=None, h=None):
    kw = {}
    if w: kw['width'] = Inches(w)
    if h: kw['height'] = Inches(h)
    return slide.shapes.add_picture(str(IMG / name), Inches(x), Inches(y), **kw)


def bar_chart(slide, x, y, w, h, title_text, labels, values):
    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series(title_text, values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = title_text
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.name = FONT
    plot = ch.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(10)
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.number_format = '0.0'
    plot.data_labels.number_format_is_linked = False
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = RGBColor(0x44, 0x44, 0x44)
    ch.category_axis.tick_labels.font.size = Pt(11)
    ch.value_axis.tick_labels.font.size = Pt(10)
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
    return ch


# ---------- 1 封面 ----------
s = new_slide()
textbox(s, 0.5, 1.8, 9, 1.2, [('撿球機', {'bullet': False})], size=48, bold=True, align=PP_ALIGN.CENTER)
textbox(s, 0.5, 3.0, 9, 0.7, [('9/15', {'bullet': False})], size=28, color=GRAY, align=PP_ALIGN.CENTER)

# ---------- 2 總覽 ----------
s = new_slide()
title(s, '7/8 至今的進度總覽 (模擬端)')
textbox(s, 0.5, 1.1, 9, 4.2, [
    '定位方案：光達 + AMCL 地圖 → 輪速編碼器 + 陀螺儀里程計 (教授建議；學長的光達是舊型室內用)',
    '弓字形巡邏 + 追球狀態機：格邊長由相機辨識距離與視角決定',
    '里程計誤差消除：找出 4 個疊加的原因，位置誤差 1.1 m → 0.03 m，航向 11° → 0.2°',
    '陀螺儀取樣：模擬 50 Hz → 1000 Hz 的原因，以及與現實 50 Hz 的差別',
    '撿球機構整合：滾輪真的把球拋進後車廂 (不再用「碰到就消除」)',
    '場地：改成開放式網球場 (無牆) + 真實尺寸網子 (擋視線、不可通過)',
    '實驗數據：隨機 10 / 15 / 20 顆球各 5 次',
], size=15, space_after=8)

# ---------- 3 定位方案 ----------
s = new_slide()
title(s, '為什麼從光達 + AMCL 改成陀螺儀里程計')
textbox(s, 0.5, 1.1, 9, 4.3, [
    '7/8 的方案 C：先驗地圖 + AMCL 動態導航 + Nav2 (Smac2D + DWB)',
    '改變原因',
    ('教授建議：網球場是開放、幾乎沒有特徵的空間，AMCL 靠雷射特徵比對校正，在空曠場地效果有限', {'indent': 1}),
    ('學長留下的 RPLIDAR A2M8 是室內用舊型光達，戶外陽光與 12 m 級的場地尺度不適用', {'indent': 1}),
    ('實體車本來就有馬達編碼器 + IMU (BNO080)，不用再多一顆感測器', {'indent': 1}),
    '新方案：航位推算 (dead reckoning)',
    ('前進速度 = 輪半徑 × 左右輪角速度平均 (編碼器)；航向 = 陀螺儀三軸角速度積分', {'indent': 1}),
    ('位置 x, y 用當下航向積分；不用 /odom、不用光達、不用地圖', {'indent': 1}),
    ('模擬優勢：Gazebo 可以直接讀真實位置，拿來跟里程計比對、量化誤差', {'indent': 1}),
], size=14)

# ---------- 4 巡邏 + 狀態機 ----------
s = new_slide()
title(s, '弓字形巡邏 + 追球狀態機')
textbox(s, 0.5, 1.1, 4.6, 4.2, [
    'PATROL：沿格子中心點走弓字形 (蛇形)，0.4 m/s，到點容差 0.3 m',
    'ALIGN：相機看到球 → 原地轉，P 控制把球對到畫面中央 (±5°)',
    'APPROACH：0.3 m/s 前進，邊走邊用像素誤差 PID 修正方向',
    'BLIND_DASH：球進相機死角後盲衝 0.3 m/s × 2 s，讓滾輪把球撈進去',
    '撿完 (球確認在後車廂) 直接回到巡邏路徑，繼續往下一格',
    '停止條件：球撿滿 / 走完路徑 / 出界 (真實座標)',
], size=13)
picture(s, 'run86_net.png', 5.2, 1.1, w=4.5)
note(s, '圖：run86，藍=PATROL、紅=APPROACH，虛線=里程計以為的位置 (幾乎完全重疊)', 5.2, 4.75, 4.5, 0.5)

# ---------- 5 格邊長 ----------
s = new_slide()
title(s, '格邊長怎麼決定')
textbox(s, 0.5, 1.1, 9, 4.3, [
    ('格邊長上限 = 2 × 0.8 × R × sin θ', {'bold': True}),
    ('R = 相機能穩定辨識球的距離 = 3.0 m (視覺節點 55 px² 面積門檻的理論距離)', {'indent': 1}),
    ('θ = 相機水平視角的一半 = 40° (視角 80°)', {'indent': 1}),
    ('0.8 = 安全係數：相鄰格子的視野重疊 20%，不漏球', {'indent': 1}),
    ('→ 上限 3.085 m；格數 = 場地邊長 ÷ 上限，無條件進位', {'indent': 1}),
    '開放場地 24 × 11 m：8 × 4 = 32 格 (每格 3.00 × 2.75 m)',
    '有網子時每個半場 11.1 × 11 m：4 × 4 = 16 格 (每格 2.775 × 2.75 m)，兩個半場共 32 格',
    '車走在格子中心線上，左右各看得到 0.8·R·sinθ = 1.54 m，蓋過半格 (1.5 m / 1.39 m)',
    '文獻依據：格子大小 = 相機 footprint × (1 − 重疊率) — Cabreira et al. 2019 (Drones) §2.2、Galceran & Carreras 2013 (RAS)；條帶間距與重疊率的計算法：Di Franco & Buttazzo 2016 (JINT)；弓字形分格：Choset & Pignon 1998',
    'footprint 寬度 = 2·R·sinθ (水平前視相機、以辨識距離 R 為半徑)；重疊率 20% → 係數 0.8',
], size=12)
s.notes_slide.notes_text_frame.text = '「2×0.8×R×sinθ」這個寫法的原始論文沒找到 (可能是崑山科大的兩篇網球撿球車碩論)；可引用的依據見 論文文獻/README.md：Cabreira 2019 §2.2 (格子大小正比於相機 footprint，解析度由重疊率決定)、Di Franco & Buttazzo 2016 (條帶間距 = footprint × (1−overlap))。'

# ---------- 6 里程計四個原因 ----------
s = new_slide()
title(s, '里程計誤差消除：四個疊加的原因')
table(s, [
    ['#', '原因', '證據', '修法', '效果'],
    ['1', 'dt 用真實時鐘，但模擬即時率 0.93', '每次 run 的 (真實/模擬時間) 都等於 (里程計/真實距離)，如 1.103 vs 1.105', 'dt 改用 /joint_states 封包的模擬時間戳 (實體車：用編碼器封包時間戳)', '距離比例 1.05 → 0.98，航向 11° → 3°'],
    ['2', '真實座標比錯點：里程計算的是輪軸中點，Gazebo 回報 base_link', '實測兩點差 (0.08, −0.135) m，原地轉時假差 0.31 m', '真實座標換算到輪軸中點', '距離比例 0.979 → 1.001'],
    ['3', '記錄時間差：兩個 timer 分開記，最多差 0.5 s', '車在動時看起來像誤差', '真實座標一到就同步記錄', '位置誤差 0.33 → 0.22 m'],
    ['4', '陀螺儀 50 Hz 取樣不夠', 'ODE 每步角速度抖動 ±8%，50 Hz 矩形積分每轉一段隨機差 ±0.5°', 'IMU 1000 Hz、在 IMU callback 用封包時間戳積分', '位置 0.026 m、航向 0.17°'],
], col_w=[0.3, 2.0, 2.6, 2.4, 1.7], row_h=0.66, size=10)
note(s, '修正後 (run74)：93.7 m 路徑，位置誤差平均 0.026 m / 最大 0.078 m，航向平均 0.17° / 最大 0.34°', y=4.75)

# ---------- 7 before/after ----------
s = new_slide()
title(s, '里程計 修正前 vs 修正後')
picture(s, 'run69_before.png', 0.3, 1.1, w=4.7)
picture(s, 'run74_after.png', 5.1, 1.1, w=4.7)
note(s, '左：run69 (修正前) 位置誤差平均 1.13 m / 最大 3.17 m / 航向 11.7°', 0.3, 4.85, 4.7, 0.5)
note(s, '右：run74 (修正後) 平均 0.026 m / 最大 0.078 m / 航向 0.17°，虛線被實線蓋住', 5.1, 4.85, 4.7, 0.5)

# ---------- 8 RTF 0.7 ----------
s = new_slide()
title(s, '單變數驗證：把模擬即時率硬壓到 0.7')
textbox(s, 0.5, 1.1, 4.4, 4.0, [
    '同一份程式、同一佈局，只把 Gazebo 物理更新率從 1000 降到 700 (即時率 0.700)',
    '同一次 run 裡並排記錄兩條里程計',
    ('新版 (封包時間戳 dt)：位置誤差平均 0.029 m / 最大 0.070 m，航向 0.15°', {'indent': 1}),
    ('舊版 (真實時鐘 dt)：路徑長 ×1.431 (=1/0.7)，位置誤差平均 18 m，航向 113°，直接飛出球場', {'indent': 1}),
    '結論：之前 7~10% 的距離與航向誤差主因就是 dt 來源，不是摩擦、碰撞或球的質量',
], size=13)
picture(s, 'run75_rtf07.png', 5.0, 1.1, w=4.7)
note(s, 'run75：黑虛線 (新) 貼著實線；洋紅虛線 (舊) 出界', 5.0, 4.9, 4.7, 0.4)

# ---------- 9 陀螺儀取樣率 ----------
s = new_slide()
title(s, '陀螺儀取樣率：模擬 50→1000 Hz 與現實 50 Hz')
textbox(s, 0.5, 1.1, 9, 4.3, [
    '模擬裡發現：Gazebo IMU 讀值跟真實角速度逐筆完全一致 (零誤差、零延遲)，但物理引擎每 1 ms 一步算出的角速度本身抖動 ±8% (接觸求解雜訊)，姿態卻是平滑的',
    '50 Hz 只取到 1/20 的步，用矩形法積分抖動訊號，每轉一段隨機差 ±0.2~0.5°；乘上 40 m 直線就是 0.5~0.9 m 的位置誤差',
    '改成 1000 Hz (= 物理步) 後每一步都算到：直走 0.003°、原地轉 246° 只差 0.22°',
    '這是模擬的數值雜訊；現實車體的角速度是連續平滑的，50 Hz 取樣不會有這種抖動',
    '現實 50 Hz 真正要注意的：(1) 用 IMU 封包自己的時間戳積分、每筆剛好算一次；(2) 三軸都要積分 (車體傾斜時只積 z 軸會少算，模擬實測少 5%)；(3) 真實陀螺儀的偏移與雜訊，模擬還沒加',
    '模擬用 1000 Hz 的目的：把積分誤差壓到只剩演算法本身，才能單獨驗證路徑與撿球邏輯',
], size=13)

# ---------- 10 三軸四元數 ----------
s = new_slide()
title(s, '航向積分改三軸四元數')
textbox(s, 0.5, 1.1, 9, 4.3, [
    '症狀：原地轉向 (ALIGN) 時偶爾一秒內少算 1~3° (實測真實轉 23°、z 軸積分只有 21.8°)',
    '原因：車體一傾斜 (滾輪擦地、被撞)，世界座標的航向變化率 = ω_z·cosφ + ω_y·sinφ + …，只積 z 軸會少算',
    '修法：三軸角速度做四元數積分 q ← q ⊗ exp(ω·dt/2)，再從 q 取 yaw',
    '單元檢查：先傾斜 20° 再繞世界 z 轉 90°，只積 z 軸得 84.6°，四元數得 90.0°',
    '效果：有網子的完整巡邏 (100 m 路徑) 航向誤差 5.3° → 0.1~0.6°，位置 1.1 m → 2~5 cm',
    '實體車 BNO080 本來就是三軸，直接套用同一套積分',
], size=14, space_after=8)

# ---------- 11 滾輪撿球 ----------
s = new_slide()
title(s, '撿球機構整合：滾輪真的把球拋進後車廂')
textbox(s, 0.5, 1.1, 5.4, 4.2, [
    '之前為了專心做路徑，撿球是「滾輪碰到球就把球消除」；現在改回物理撿球',
    '關鍵是滾輪位置：5/21 影片能成功的設定 (5/13 筆記) 是滾輪在 x=0.2、z=0.03；後來改回 CAD 位置 (x=0.10) 後離後車廂 17.7 cm 前擋牆只剩 10 cm，球被推到牆邊只能垂直彈起、進不了車廂 (已跟設計組說明)',
    '滾輪常轉，轉速掃描 (右表)：W=40~45 (62~70 rad/s) 在盲衝 0.2~0.4 m/s 全部進車廂；≥50 拋過頭飛出車尾；採用 W=42.5 (約 66 rad/s)',
    '「撿到」判定改用真實座標：球躺在車廂地板高度 (z 0.05~0.15) 且連續 3 次都在車廂內；結束時稽核每顆球位置',
], size=12)
table(s, [
    ['W', 'rad/s', 'v=0.2', 'v=0.3', 'v=0.4'],
    ['35', '54', '進', '未過牆', '未過牆'],
    ['40', '62', '進', '進', '進'],
    ['45', '70', '進', '進', '進'],
    ['50', '78', '飛出', '飛出', '飛出'],
    ['55', '86', '—', '飛出', '飛出'],
], x=6.1, y=1.1, w=3.55, col_w=[0.6, 0.7, 0.75, 0.75, 0.75], row_h=0.3, size=11)
picture(s, 'video47_pickup.jpg', 6.1, 3.15, w=3.55)
note(s, '5/21 影片：7.5 s 球到滾輪，7.75 s 飛過車廂前緣', 6.1, 4.3, 3.55, 0.4)

# ---------- 12 撿球結果 ----------
s = new_slide()
title(s, '物理撿球 + 巡邏：完整測試')
textbox(s, 0.5, 1.1, 4.6, 4.2, [
    '標準 9 顆球佈局：run76、run77 皆 9/9 進後車廂，結束稽核 9 顆全部躺在車廂地板 (一顆疊在第二層)',
    '後車廂載球對里程計的影響 (實測)：位置誤差 0.026 → 0.056~0.078 m，航向 0.17° → 0.33~0.63°；在前 60 秒追球期間累積後持平，撿球瞬間本身沒有跳動',
    '設計組：後車廂可放 20 顆；本次最多測 20 顆',
    '注意：滾輪轉速 ≥ 90 rad/s 時球卡住會讓滾輪關節在物理引擎裡失控，轉速保持在 66 rad/s 附近',
], size=13)
picture(s, 'run77_pickup.png', 5.2, 1.1, w=4.5)
note(s, 'run77：9/9，位置誤差平均 0.078 m', 5.2, 4.75, 4.5, 0.4)

# ---------- 13 開放場地 ----------
s = new_slide()
title(s, '場地改成開放式網球場 (無牆)')
textbox(s, 0.5, 1.1, 4.6, 4.2, [
    '現實場地是開放的：Gazebo 裡的牆只留半透明視覺當邊界，碰撞全部拿掉',
    '「撞牆停止」改成「出界停止」：真實座標離開場地邊界外緣就停 (模擬才有的判定)',
    '巡邏範圍改成整個 24 × 11 m 雙打場地 (不再內縮 1 m)，格數由格邊長公式算：8 × 4 = 32 格',
    '隨機 9 顆球 × 3 次 (run78~80)：3/3 皆 9/9，稽核全部在車廂',
], size=13)
picture(s, 'run80_open.png', 5.2, 1.1, w=4.5)
note(s, 'run80：開放場地隨機佈局 9/9', 5.2, 4.75, 4.5, 0.4)

# ---------- 14 網子 ----------
s = new_slide()
title(s, '真實網子：擋視線、不可通過')
textbox(s, 0.5, 1.1, 5.0, 4.3, [
    'ITF 規格：中央 0.914 m、網柱 1.07 m，網柱在雙打邊線外 0.914 m → 全長 12.8 m (網柱 y = ±6.4)',
    '模擬：x=0 處 1 m 高的深灰實心板 + 兩根網柱，有碰撞；從 15 cm 高的相機看過去完全擋住另一半場 (對應 YOLO 隔網辨識不可靠)',
    '路徑：Choset boustrophedon cellular decomposition 的最簡單特例——網子把場地切成兩個矩形，各自走弓字 (每半場 4 × 4 格，從離網 0.9 m 起)，第一半場走完繞網柱外側 (y=±7.5) 到第二半場靠網的角落接上',
    '追球時的網子保護：車體中心離網 < 0.763 m (車身 0.5 m + 車頭 0.263 m) 又朝向網子就放棄這顆球，5 秒不理視覺',
    '隨機 9 顆球 (離網 ≥ 1 m) × 3 次 (run85~87)：3/3 皆 9/9，稽核全在車廂，零次觸發放棄；位置誤差 2~5 cm、航向 ≤ 0.6°',
], size=12)
picture(s, 'run86_net.png', 5.6, 1.1, w=4.1)
note(s, 'run86：兩個半場 + 繞網柱', 5.6, 4.5, 4.1, 0.4)

# ---------- 15/16 實驗數據 ----------
data_path = HERE / 'data_0915.json'
s = new_slide()
title(s, '實驗數據：隨機 10 / 15 / 20 顆球 × 5 次')
if data_path.exists():
    d = json.loads(data_path.read_text())
    rows = [['球數', '成功次數', '平均撿完時間 (s)', '平均每顆 (s)', '平均路徑 (m)', '位置誤差平均 (m)', '最終偏移 (m)', '航向誤差平均 (°)']]
    for g in d['groups']:
        rows.append([g['balls'], f"{g['success']}/{g['runs']}", g['time'], g['sec_per_ball'], g['path'], g['pos_err'], g['pos_final'], g['yaw_err']])
    table(s, rows, col_w=[0.8, 1.0, 1.4, 1.1, 1.2, 1.4, 1.0, 1.1], row_h=0.4, size=11)
    labels = [f"{g['balls']} 顆" for g in d['groups']]
    bar_chart(s, 0.5, 2.9, 4.4, 2.3, '平均撿完時間 (s)', labels, [g['time'] for g in d['groups']])
    bar_chart(s, 5.2, 2.9, 4.4, 2.3, '里程計最終偏移 (m)', labels, [g['pos_final'] for g in d['groups']])
    if d.get('note'):
        note(s, d['note'], y=5.25, h=0.3)

    s = new_slide()
    title(s, '每次 run 明細')
    rows = [['球數', 'run', '撿到', '撿完時間 (s)', '路徑 (m)', '位置誤差 平均/最大 (m)', '最終偏移 (m)', '航向 平均/最大 (°)']]
    for r in d['runs']:
        rows.append([r['balls'], r['run'], f"{r['picked']}/{r['balls']}", r['time'], r['path'], f"{r['pos_mean']} / {r['pos_max']}", r['pos_final'], f"{r['yaw_mean']} / {r['yaw_max']}"])
    table(s, rows, col_w=[0.6, 0.8, 0.7, 1.1, 0.9, 1.9, 1.0, 1.6], row_h=0.245, size=9)
else:
    textbox(s, 0.5, 1.1, 9, 2, ['(數據跑完後由 summarize_runs.py 產生 data_0915.json，再重新執行 make_0915.py)'], size=14)

# ---------- 數據觀察 ----------
s = new_slide()
title(s, '實驗數據觀察')
textbox(s, 0.5, 1.1, 9, 4.3, [
    '10 顆：5/5 全撿滿，平均 369 s (每顆 37 s)；15 顆：4/5，平均 470 s (每顆 31 s)；20 顆：1/5，撿滿那次 492 s (每顆 25 s)',
    '球越多、每顆平均時間越短：巡邏路徑固定 (兩個半場 32 格)，球多時追球距離短、巡邏空跑的比例低',
    '里程計：13/15 次位置誤差平均 < 0.15 m、最終偏移 < 0.2 m、航向 < 1°，跟 9 顆的結果一致，載 15~20 顆球沒有讓里程計變差',
    '失敗模式 (誠實列出)：',
    ('run103 (14/15)、run109 (19/20)：稽核發現漏掉的那顆球在場外 36 m / 289 m 遠——球被滾輪或車廂內的球碰撞後高速彈飛 (物理模擬的問題，現實不會)', {'indent': 1}),
    ('run110、run111 (19/20)：各漏 1 顆在場內 (x=-7.3 / x=+1.6)，原因未逐一分析 (run110 有 2 次靠網放棄追球)', {'indent': 1}),
    ('run108 (8/20)：t=227 s 時 /joint_states 訂閱斷掉、里程計凍結，車直線出界；同樣症狀之前出現過，跟 Gazebo/DDS 負載有關 (每 tick 對每顆球查一次位置時 15 顆就必發生，改成輪流查一顆後 10 次只剩這 1 次)', {'indent': 1}),
    '20 顆的「最終偏移 0.8 m」平均是被 run108 的 3.07 m 拉高的，其餘 4 次都 < 0.8 m、3 次 < 0.1 m',
], size=12)

# ---------- 17 限制與後續 ----------
s = new_slide()
title(s, '已知限制與後續')
textbox(s, 0.5, 1.1, 9, 4.3, [
    '模擬與實體的差異：陀螺儀還是理想 sensor (沒有偏移、雜訊)；滾輪模型位置與 CAD 不同 (設計組已知)；前車體無碰撞',
    '實體車要對應的三件事：編碼器 / IMU 封包時間戳、三軸陀螺儀四元數積分、滾輪轉速約 66 rad/s',
    '格邊長公式的文獻出處待補',
    '後續：加入陀螺儀偏移 / 雜訊模型驗證現實可行性；多顆球載重下的行為；與實體車比對',
], size=14, space_after=8)

# ---------- 未來工作：實體車轉向對準 ----------
s = new_slide()
title(s, '未來工作：把模擬的對準 PID 搬到實體車')
picture(s, 'future_a_stairs.png', 0.4, 1.05, w=4.6)
picture(s, 'future_b_divergent.png', 5.0, 1.05, w=4.6)
textbox(s, 0.4, 3.75, 9.2, 1.7, [
    '現況 (B 同學的實體車做法，圖 a)：每次只轉 5° 就停 1 秒等相機更新再判斷，階梯式收斂，35° 要 8 秒以上',
    '風險 (圖 b)：直接套模擬的增益、不管深度相機更新延遲，會轉過頭、來回擺盪甚至發散',
    '目標：模擬用的對準 P / 接近 PID 同一套程式直接上實體車，只調參數 (Kp、Kd、角速度上限、取樣週期) 去吸收深度相機的更新延遲，得到穩定收斂的 step response (灰虛線)',
    '為什麼是 step response：對準時車是靜止只轉向，球角度從 ±40° 一步到 0° 就是標準的階躍響應，可以直接量上升時間、超調、穩定時間來比較參數',
], size=12, space_after=4)
note(s, '兩張圖為示意 (相機延遲設 0.4 s)，不是量測數據；x 軸時間、y 軸球相對車體中心的夾角 ±40°', y=5.3, h=0.3)

out = HERE / '撿球機9_15.pptx'
prs.save(str(out))
print('wrote', out)
